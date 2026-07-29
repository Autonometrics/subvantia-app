"""
Subvantia - aplicacion web funcional (prototipo real, no mockup).

Ejecutar:
    pip install -r requirements.txt
    python app.py
Luego abrir http://localhost:5000  (usuario demo: demo@subvantia.test / subvantia2026)
"""

import io
import os
import re
import datetime as dt

from flask import (
    Flask, render_template, request, redirect, url_for, session, flash,
    send_file, jsonify, g,
)
import secrets
import string
from werkzeug.security import check_password_hash, generate_password_hash
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import database as db
import matching

app = Flask(__name__)
app.secret_key = os.environ.get("SUBVANTIA_SECRET", "clave-de-desarrollo-cambiar-en-produccion")

# La cuenta demo es publica y compartida por cualquier visitante que reciba el
# enlace: varias gestorias distintas pueden estar mirandola a la vez. Por eso
# es de solo consulta -- no puede importar carteras propias, mover expedientes
# ni cambiar su contrasena compartida, para que un visitante no altere lo que
# esta viendo otro. Quien quiera probarla con sus propios datos necesita una
# cuenta individual (la crea el admin desde "Gestorias").
EMAIL_CUENTA_DEMO = "demo@subvantia.test"
MENSAJE_DEMO_SOLO_LECTURA = (
    "Esta es la cuenta de demostración pública: es solo para ver cómo "
    "funciona Subvantia con datos de ejemplo, no para modificar nada (la "
    "comparten varias gestorías a la vez). Si quieres probarlo con tu propia "
    "cartera de clientes, pide tu acceso a Subvantia y te creamos una cuenta "
    "con tu email."
)


def es_cuenta_demo():
    return session.get("usuario_email") == EMAIL_CUENTA_DEMO


# ---------------------------------------------------------------- utilidades

def get_conn():
    if "db_conn" not in g:
        g.db_conn = db.get_db()
    return g.db_conn


@app.teardown_appcontext
def cerrar_conn(exception=None):
    conn = g.pop("db_conn", None)
    if conn is not None:
        conn.close()


def login_requerido(view):
    def wrapped(*args, **kwargs):
        if not session.get("usuario_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    wrapped.__name__ = view.__name__
    return wrapped


def admin_requerido(view):
    """Como login_requerido, pero ademas exige rol de administrador. Anadir o
    editar convocatorias es tarea de Subvantia (mantener la base de datos al
    dia es el motivo por el que las gestorias pagan), no de las gestorias
    clientes: ellas solo consultan la base, no la editan."""
    def wrapped(*args, **kwargs):
        if not session.get("usuario_id"):
            return redirect(url_for("login"))
        if session.get("usuario_rol") != "admin":
            flash("Solo un administrador de Subvantia puede añadir o editar convocatorias.")
            return redirect(url_for("convocatorias"))
        return view(*args, **kwargs)
    wrapped.__name__ = view.__name__
    return wrapped


# --------------------------------------------------------------------- auth

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        conn = get_conn()
        usuario = conn.execute("SELECT * FROM usuarios WHERE email = ?", (email,)).fetchone()
        if usuario and check_password_hash(usuario["password_hash"], password):
            rol = usuario["rol"] or "gestoria"
            estado = usuario["estado_suscripcion"] or "activa"
            # El admin (Subvantia) nunca queda bloqueado por estado de suscripcion,
            # ese concepto solo aplica a las cuentas de gestoria clientes.
            if rol != "admin" and estado != "activa":
                mensajes_estado = {
                    "impago": "Esta cuenta está suspendida por impago. Contacta con Subvantia para reactivarla.",
                    "pausada": "Esta cuenta está pausada. Contacta con Subvantia para reactivarla.",
                }
                flash(mensajes_estado.get(estado, "Esta cuenta no está activa. Contacta con Subvantia."))
                return render_template("login.html", mostrar_pistas_demo=not db.USANDO_POSTGRES)
            session["usuario_id"] = usuario["id"]
            session["usuario_nombre"] = usuario["nombre"]
            session["usuario_gestoria"] = usuario["gestoria"]
            session["usuario_rol"] = rol
            session["usuario_email"] = usuario["email"]
            ip_origen = request.headers.get("X-Forwarded-For", request.remote_addr or "")
            ip_origen = ip_origen.split(",")[0].strip()
            conn.execute(
                "INSERT INTO accesos_log (usuario_id, fecha_hora, ip) VALUES (?, ?, ?)",
                (usuario["id"], dt.datetime.now().isoformat(), ip_origen),
            )
            conn.commit()
            return redirect(url_for("dashboard"))
        flash("Email o contraseña incorrectos.")
    # Las credenciales de demo/admin en pantalla son solo una ayuda para
    # desarrollo local (sqlite, sin DATABASE_URL). En cuanto la app corre con
    # una base de datos real (Postgres/Neon, es decir un despliegue de
    # verdad), se ocultan automaticamente para que ninguna visita publica a la
    # web pueda leer la contraseña de administrador.
    return render_template("login.html", mostrar_pistas_demo=not db.USANDO_POSTGRES)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/cambiar-password", methods=["GET", "POST"])
@login_requerido
def cambiar_password():
    if request.method == "POST":
        if es_cuenta_demo():
            flash(MENSAJE_DEMO_SOLO_LECTURA)
            return redirect(url_for("dashboard"))
        actual = request.form.get("actual", "")
        nueva = request.form.get("nueva", "")
        confirmar = request.form.get("confirmar", "")

        conn = get_conn()
        usuario = conn.execute(
            "SELECT * FROM usuarios WHERE id = ?", (session["usuario_id"],)
        ).fetchone()

        if not check_password_hash(usuario["password_hash"], actual):
            flash("La contraseña actual no es correcta.")
        elif len(nueva) < 6:
            flash("La nueva contraseña debe tener al menos 6 caracteres.")
        elif nueva != confirmar:
            flash("La nueva contraseña y su confirmación no coinciden.")
        else:
            conn.execute(
                "UPDATE usuarios SET password_hash = ? WHERE id = ?",
                (generate_password_hash(nueva), session["usuario_id"]),
            )
            conn.commit()
            flash("Contraseña actualizada correctamente.")
            return redirect(url_for("dashboard"))

    return render_template("cambiar_password.html")


# ---------------------------------------------------------------- dashboard

@app.route("/")
def home():
    # Si ya hay sesion, directo al dashboard como siempre. Si no, en vez de
    # forzar el formulario de login a pelo, se muestra una pagina publica
    # breve (quien es Subvantia, que hace) — asi la URL real del producto
    # tiene una cara presentable para cualquiera que la reciba antes de tener
    # cuenta, en vez de un formulario de login sin contexto.
    if session.get("usuario_id"):
        return redirect(url_for("dashboard"))
    return render_template("home.html")


@app.route("/dashboard")
@login_requerido
def dashboard():
    conn = get_conn()
    uid = session["usuario_id"]

    total_clientes = conn.execute(
        "SELECT COUNT(*) AS n FROM clientes WHERE usuario_id = ?", (uid,)
    ).fetchone()["n"]

    total_matches = conn.execute(
        "SELECT COUNT(*) AS n FROM matches WHERE usuario_id = ?", (uid,)
    ).fetchone()["n"]

    # Desde que el motor de cruce (matching.py) exige CCAA, CNAE, forma
    # juridica y antiguedad como requisitos excluyentes, todo match que
    # sobrevive cumple el 100% de los criterios objetivos: la banda
    # alta/media/baja ya no distingue nada (todo cae en "alta"). Por eso el
    # dashboard ya no desglosa por banda -- en su lugar avisa de lo que si
    # varia entre matches: cuanta urgencia tiene (fecha de cierre) y cuanto
    # dinero hay en juego (cuantia).
    dinero_potencial = conn.execute(
        "SELECT COALESCE(SUM(cuantia_estimada), 0) AS total FROM matches WHERE usuario_id = ?",
        (uid,),
    ).fetchone()["total"]

    convocatorias_activas = conn.execute(
        "SELECT COUNT(*) AS n FROM convocatorias WHERE estado_convocatoria IN ('abierta','verificar_urgente')"
    ).fetchone()["n"]

    alertas = conn.execute(
        """SELECT DISTINCT conv.id_convocatoria, conv.titulo, conv.estado_convocatoria, conv.fecha_cierre
           FROM matches m JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ? AND conv.estado_convocatoria = 'verificar_urgente'""",
        (uid,),
    ).fetchall()

    top_convocatorias = conn.execute(
        """SELECT conv.id_convocatoria, conv.titulo, conv.estado_convocatoria, conv.fecha_cierre,
                  COUNT(DISTINCT m.cliente_id) AS clientes_elegibles
           FROM matches m JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ?
           GROUP BY conv.id_convocatoria
           ORDER BY clientes_elegibles DESC
           LIMIT 6""",
        (uid,),
    ).fetchall()

    # Cierran pronto: convocatorias con match, abiertas, con fecha de cierre
    # confirmada dentro de los proximos 30 dias -- para priorizar a que
    # clientes contactar primero.
    hoy = dt.date.today().isoformat()
    limite_urgencia = (dt.date.today() + dt.timedelta(days=30)).isoformat()
    cierran_pronto = conn.execute(
        """SELECT conv.id_convocatoria, conv.titulo, conv.fecha_cierre,
                  COUNT(DISTINCT m.cliente_id) AS clientes_elegibles
           FROM matches m JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ? AND conv.estado_convocatoria IN ('abierta', 'verificar_urgente')
                 AND conv.fecha_cierre IS NOT NULL AND conv.fecha_cierre <> ''
                 AND conv.fecha_cierre BETWEEN ? AND ?
           GROUP BY conv.id_convocatoria
           ORDER BY conv.fecha_cierre ASC
           LIMIT 8""",
        (uid, hoy, limite_urgencia),
    ).fetchall()
    cierran_pronto = [
        {
            "id_convocatoria": r["id_convocatoria"],
            "titulo": r["titulo"],
            "fecha_cierre": r["fecha_cierre"],
            "clientes_elegibles": r["clientes_elegibles"],
            "dias_restantes": (
                dt.datetime.strptime(str(r["fecha_cierre"])[:10], "%Y-%m-%d").date() - dt.date.today()
            ).days,
        }
        for r in cierran_pronto
    ]

    # Mayor cuantia potencial: convocatorias con match ordenadas por cuantia
    # maxima descendente -- para priorizar por impacto economico.
    mayor_cuantia = conn.execute(
        """SELECT conv.id_convocatoria, conv.titulo, conv.cuantia_max_eur,
                  COUNT(DISTINCT m.cliente_id) AS clientes_elegibles
           FROM matches m JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ? AND conv.cuantia_max_eur IS NOT NULL
           GROUP BY conv.id_convocatoria
           ORDER BY conv.cuantia_max_eur DESC
           LIMIT 6""",
        (uid,),
    ).fetchall()

    return render_template(
        "dashboard.html",
        total_clientes=total_clientes,
        total_matches=total_matches,
        dinero_potencial=dinero_potencial,
        convocatorias_activas=convocatorias_activas,
        alertas=alertas,
        top_convocatorias=top_convocatorias,
        cierran_pronto=cierran_pronto,
        mayor_cuantia=mayor_cuantia,
    )


# ----------------------------------------------------------------- clientes

@app.route("/clientes")
@login_requerido
def clientes():
    conn = get_conn()
    uid = session["usuario_id"]
    filas = conn.execute(
        "SELECT * FROM clientes WHERE usuario_id = ? ORDER BY nombre", (uid,)
    ).fetchall()
    return render_template("clientes.html", clientes=filas)


@app.route("/clientes/importar", methods=["POST"])
@login_requerido
def importar_clientes():
    if es_cuenta_demo():
        flash(MENSAJE_DEMO_SOLO_LECTURA)
        return redirect(url_for("clientes"))
    conn = get_conn()
    uid = session["usuario_id"]
    archivo = request.files.get("archivo")
    if not archivo or archivo.filename == "":
        flash("Selecciona un fichero Excel para importar.")
        return redirect(url_for("clientes"))

    try:
        wb = openpyxl.load_workbook(archivo, data_only=True)
        hoja = wb["Cartera"] if "Cartera" in wb.sheetnames else wb.active
        headers = [c.value for c in hoja[1]]
        idx = {h: i for i, h in enumerate(headers) if h}

        columnas_necesarias = ["Nombre", "CNAE", "Comunidad Autonoma", "Forma juridica", "Fecha de alta"]
        faltan = [c for c in columnas_necesarias if c not in idx]
        if faltan:
            flash(f"Faltan columnas obligatorias en el Excel: {', '.join(faltan)}")
            return redirect(url_for("clientes"))

        insertados = 0
        nuevos_ids = []
        for row in hoja.iter_rows(min_row=2, values_only=True):
            if row[idx["Nombre"]] is None:
                continue
            fecha_alta = row[idx["Fecha de alta"]]
            if hasattr(fecha_alta, "strftime"):
                fecha_alta = fecha_alta.strftime("%Y-%m-%d")
            else:
                fecha_alta = str(fecha_alta)[:10] if fecha_alta else None

            # La plantilla ofrece el CNAE como desplegable "62 - Programacion,
            # consultoria..." para que no haya que memorizar codigos; aqui nos
            # quedamos solo con el codigo numerico inicial antes de guardarlo,
            # para que la tabla de clientes y las exportaciones muestren un
            # dato limpio en vez del texto completo del desplegable. Sigue
            # aceptando el formato antiguo (solo el codigo, sin descripcion).
            cnae_valor = row[idx["CNAE"]]
            m = re.match(r"^\s*(\d+)", str(cnae_valor)) if cnae_valor is not None else None
            cnae_codigo = m.group(1) if m else cnae_valor

            sql_insertar_cliente = """INSERT INTO clientes
                   (usuario_id, cif_nif, nombre, cnae, comunidad_autonoma, provincia,
                    forma_juridica, fecha_alta, gestor_asignado, activo, creado_en)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)"""
            if db.USANDO_POSTGRES:
                # Postgres no expone cursor.lastrowid; lo obtenemos pidiendo
                # explicitamente el id generado con RETURNING (ver database.py).
                sql_insertar_cliente += " RETURNING id"
            cur = conn.execute(
                sql_insertar_cliente,
                (
                    uid,
                    row[idx.get("CIF/NIF")] if "CIF/NIF" in idx else None,
                    row[idx["Nombre"]],
                    cnae_codigo,
                    row[idx["Comunidad Autonoma"]],
                    row[idx.get("Provincia")] if "Provincia" in idx else None,
                    row[idx["Forma juridica"]],
                    fecha_alta,
                    row[idx.get("Gestor asignado")] if "Gestor asignado" in idx else None,
                    dt.datetime.now().isoformat(),
                ),
            )
            nuevos_ids.append(cur.lastrowid)
            insertados += 1
        conn.commit()

        total_matches = 0
        for cid in nuevos_ids:
            total_matches += matching.recalcular_matches_cliente(conn, cid, uid)

        flash(f"Importados {insertados} clientes. Generados {total_matches} matches automaticamente.")
    except Exception as e:
        flash(f"Error al importar el Excel: {e}")

    return redirect(url_for("clientes"))


# Listas cerradas para los desplegables de la plantilla de clientes. Los
# mismos nombres canonicos que ya reconoce matching.py (_CCAA_ALIASES /
# _clasificar_forma_juridica), para que un dato elegido del desplegable
# SIEMPRE se reconozca en el cruce, sin depender de la normalizacion de
# texto libre -- el desplegable evita el error de raiz, no lo corrige a
# posteriori.
CCAA_LISTA = [
    "Andalucía", "Aragón", "Asturias", "Baleares", "Canarias", "Cantabria",
    "Castilla y León", "Castilla-La Mancha", "Cataluña", "Comunidad Valenciana",
    "Extremadura", "Galicia", "Madrid", "Murcia", "Navarra", "País Vasco",
    "La Rioja", "Ceuta", "Melilla",
]

FORMA_JURIDICA_LISTA = ["Autónomo", "SL", "SA", "Otra"]

# CNAE-2009 (INE): las 88 divisiones oficiales de 2 digitos. Se muestran como
# "codigo - descripcion" para que quien rellena el Excel elija el sector por
# su nombre en vez de tener que saberse el codigo de memoria. El motor de
# cruce (matching.py) ya solo mira los 2 primeros caracteres del campo CNAE,
# asi que basta con quedarnos con el codigo al importar (ver mas abajo).
CNAE_DIVISIONES = [
    ("01", "Agricultura, ganadería, caza y servicios relacionados"),
    ("02", "Silvicultura y explotación forestal"),
    ("03", "Pesca y acuicultura"),
    ("05", "Extracción de antracita, hulla y lignito"),
    ("06", "Extracción de crudo de petróleo y gas natural"),
    ("07", "Extracción de minerales metálicos"),
    ("08", "Otras industrias extractivas"),
    ("09", "Actividades de apoyo a las industrias extractivas"),
    ("10", "Industria de la alimentación"),
    ("11", "Fabricación de bebidas"),
    ("12", "Industria del tabaco"),
    ("13", "Industria textil"),
    ("14", "Confección de prendas de vestir"),
    ("15", "Industria del cuero y del calzado"),
    ("16", "Industria de la madera y del corcho, excepto muebles"),
    ("17", "Industria del papel"),
    ("18", "Artes gráficas y reproducción de soportes grabados"),
    ("19", "Coquerías y refino de petróleo"),
    ("20", "Industria química"),
    ("21", "Fabricación de productos farmacéuticos"),
    ("22", "Fabricación de productos de caucho y plásticos"),
    ("23", "Fabricación de otros productos minerales no metálicos"),
    ("24", "Metalurgia; fabricación de productos de hierro, acero y ferroaleaciones"),
    ("25", "Fabricación de productos metálicos, excepto maquinaria y equipo"),
    ("26", "Fabricación de productos informáticos, electrónicos y ópticos"),
    ("27", "Fabricación de material y equipo eléctrico"),
    ("28", "Fabricación de maquinaria y equipo n.c.o.p."),
    ("29", "Fabricación de vehículos de motor, remolques y semirremolques"),
    ("30", "Fabricación de otro material de transporte"),
    ("31", "Fabricación de muebles"),
    ("32", "Otras industrias manufactureras"),
    ("33", "Reparación e instalación de maquinaria y equipo"),
    ("35", "Suministro de energía eléctrica, gas, vapor y aire acondicionado"),
    ("36", "Captación, depuración y distribución de agua"),
    ("37", "Recogida y tratamiento de aguas residuales"),
    ("38", "Recogida, tratamiento y eliminación de residuos; valorización"),
    ("39", "Actividades de descontaminación y otros servicios de gestión de residuos"),
    ("41", "Construcción de edificios"),
    ("42", "Ingeniería civil"),
    ("43", "Actividades de construcción especializada"),
    ("45", "Venta y reparación de vehículos de motor y motocicletas"),
    ("46", "Comercio al por mayor e intermediarios del comercio"),
    ("47", "Comercio al por menor"),
    ("49", "Transporte terrestre y por tubería"),
    ("50", "Transporte marítimo y por vías navegables interiores"),
    ("51", "Transporte aéreo"),
    ("52", "Almacenamiento y actividades anexas al transporte"),
    ("53", "Actividades postales y de correos"),
    ("55", "Servicios de alojamiento"),
    ("56", "Servicios de comidas y bebidas"),
    ("58", "Edición"),
    ("59", "Actividades cinematográficas, de vídeo, de programas de TV y grabación musical"),
    ("60", "Actividades de programación y emisión de radio y televisión"),
    ("61", "Telecomunicaciones"),
    ("62", "Programación, consultoría y otras actividades relacionadas con la informática"),
    ("63", "Servicios de información"),
    ("64", "Servicios financieros, excepto seguros y fondos de pensiones"),
    ("65", "Seguros, reaseguros y fondos de pensiones, excepto Seguridad Social"),
    ("66", "Actividades auxiliares a los servicios financieros y a los seguros"),
    ("68", "Actividades inmobiliarias"),
    ("69", "Actividades jurídicas y de contabilidad"),
    ("70", "Actividades de las sedes centrales; consultoría de gestión empresarial"),
    ("71", "Servicios técnicos de arquitectura e ingeniería; ensayos y análisis"),
    ("72", "Investigación y desarrollo"),
    ("73", "Publicidad y estudios de mercado"),
    ("74", "Otras actividades profesionales, científicas y técnicas"),
    ("75", "Actividades veterinarias"),
    ("77", "Actividades de alquiler"),
    ("78", "Actividades relacionadas con el empleo"),
    ("79", "Agencias de viajes, operadores turísticos y reservas"),
    ("80", "Actividades de seguridad e investigación"),
    ("81", "Servicios a edificios y actividades de jardinería"),
    ("82", "Actividades administrativas de oficina y otras auxiliares a empresas"),
    ("84", "Administración Pública y defensa; Seguridad Social obligatoria"),
    ("85", "Educación"),
    ("86", "Actividades sanitarias"),
    ("87", "Asistencia en establecimientos residenciales"),
    ("88", "Actividades de servicios sociales sin alojamiento"),
    ("90", "Actividades de creación, artísticas y espectáculos"),
    ("91", "Bibliotecas, archivos, museos y otras actividades culturales"),
    ("92", "Actividades de juegos de azar y apuestas"),
    ("93", "Actividades deportivas, recreativas y de entretenimiento"),
    ("94", "Actividades asociativas"),
    ("95", "Reparación de ordenadores, efectos personales y artículos de uso doméstico"),
    ("96", "Otros servicios personales"),
    ("97", "Actividades de los hogares como empleadores de personal doméstico"),
    ("98", "Actividades de los hogares como productores de bienes y servicios para uso propio"),
    ("99", "Actividades de organizaciones y organismos extraterritoriales"),
]
CNAE_LISTA = [f"{codigo} - {desc}" for codigo, desc in CNAE_DIVISIONES]


@app.route("/clientes/plantilla")
@login_requerido
def descargar_plantilla():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cartera"
    headers = ["CIF/NIF", "Nombre", "CNAE", "Comunidad Autonoma", "Provincia",
               "Forma juridica", "Fecha de alta", "Gestor asignado"]
    ws.append(headers)
    ws.append(["B00000000", "Empresa Ejemplo SL", "62 - Programación, consultoría y otras actividades relacionadas con la informática",
                "Madrid", "Madrid", "SL", "2019-03-01", "Ana Garcia"])
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.column_dimensions["C"].width = 55  # CNAE: cabe "codigo - descripcion" entero

    # Hoja auxiliar oculta con las listas cerradas -- los desplegables de la
    # hoja "Cartera" apuntan aqui en vez de llevar la lista incrustada, para
    # no toparse con el limite de ~255 caracteres de una lista en linea.
    ws_listas = wb.create_sheet("Listas")
    ws_listas["A1"] = "Comunidad Autonoma"
    ws_listas["B1"] = "Forma juridica"
    ws_listas["C1"] = "CNAE"
    for i, v in enumerate(CCAA_LISTA, start=2):
        ws_listas.cell(row=i, column=1, value=v)
    for i, v in enumerate(FORMA_JURIDICA_LISTA, start=2):
        ws_listas.cell(row=i, column=2, value=v)
    for i, v in enumerate(CNAE_LISTA, start=2):
        ws_listas.cell(row=i, column=3, value=v)
    ws_listas.sheet_state = "hidden"

    n_filas = 1000  # margen de sobra para cualquier cartera real
    dv_ccaa = DataValidation(type="list", formula1=f"Listas!$A$2:$A${len(CCAA_LISTA)+1}",
                              allow_blank=True, showErrorMessage=True, errorStyle="stop",
                              errorTitle="Comunidad no reconocida",
                              error="Elige una comunidad autónoma de la lista desplegable.",
                              showInputMessage=True, promptTitle="Comunidad autónoma",
                              prompt="Haz clic en la flecha y elige de la lista.")
    dv_fj = DataValidation(type="list", formula1=f"Listas!$B$2:$B${len(FORMA_JURIDICA_LISTA)+1}",
                             allow_blank=True, showErrorMessage=True, errorStyle="stop",
                             errorTitle="Forma jurídica no reconocida",
                             error="Elige una forma jurídica de la lista desplegable.",
                             showInputMessage=True, promptTitle="Forma jurídica",
                             prompt="Haz clic en la flecha y elige de la lista.")
    dv_cnae = DataValidation(type="list", formula1=f"Listas!$C$2:$C${len(CNAE_LISTA)+1}",
                               allow_blank=True, showErrorMessage=True, errorStyle="stop",
                               errorTitle="Sector no reconocido",
                               error="Elige el sector de la lista desplegable (empieza a escribir para buscar).",
                               showInputMessage=True, promptTitle="Sector de actividad (CNAE)",
                               prompt="Haz clic en la flecha y elige el sector. Con el desplegable abierto, "
                                      "escribir letras salta a la primera coincidencia.")
    ws.add_data_validation(dv_ccaa)
    ws.add_data_validation(dv_fj)
    ws.add_data_validation(dv_cnae)
    dv_ccaa.add(f"D2:D{n_filas}")
    dv_fj.add(f"F2:F{n_filas}")
    dv_cnae.add(f"C2:C{n_filas}")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="Plantilla_Importacion_Clientes.xlsx",
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------ convocatorias

@app.route("/convocatorias")
@login_requerido
def convocatorias():
    conn = get_conn()
    filas = conn.execute("SELECT * FROM convocatorias ORDER BY estado_convocatoria, fecha_cierre").fetchall()
    return render_template("convocatorias.html", convocatorias=filas)


CAMPOS_CONVOCATORIA = [
    "titulo", "origen", "comunidad_autonoma", "cnae_elegibles", "forma_juridica_elegible",
    "antiguedad_min_meses", "antiguedad_max_meses", "cuantia_min_eur", "cuantia_max_eur",
    "fecha_apertura", "fecha_cierre", "requisitos_texto", "url_fuente",
    "estado_convocatoria", "verificado_manualmente",
]
CAMPOS_NUMERICOS = {"antiguedad_min_meses", "antiguedad_max_meses", "cuantia_min_eur", "cuantia_max_eur"}


def _leer_formulario_convocatoria():
    datos = {}
    for campo in CAMPOS_CONVOCATORIA:
        valor = request.form.get(campo, "").strip()
        if campo in CAMPOS_NUMERICOS:
            datos[campo] = float(valor) if valor != "" else None
        else:
            datos[campo] = valor if valor != "" else None
    return datos


@app.route("/convocatorias/nueva", methods=["GET", "POST"])
@admin_requerido
def nueva_convocatoria():
    conn = get_conn()
    if request.method == "POST":
        id_convocatoria = request.form.get("id_convocatoria", "").strip()
        if not id_convocatoria:
            flash("El ID de la convocatoria es obligatorio.")
            return redirect(url_for("nueva_convocatoria"))
        existe = conn.execute(
            "SELECT 1 FROM convocatorias WHERE id_convocatoria = ?", (id_convocatoria,)
        ).fetchone()
        if existe:
            flash(f"Ya existe una convocatoria con el ID {id_convocatoria}.")
            return redirect(url_for("nueva_convocatoria"))

        datos = _leer_formulario_convocatoria()
        columnas = ["id_convocatoria"] + CAMPOS_CONVOCATORIA
        valores = [id_convocatoria] + [datos[c] for c in CAMPOS_CONVOCATORIA]
        placeholders = ", ".join("?" for _ in columnas)
        conn.execute(
            f"INSERT INTO convocatorias ({', '.join(columnas)}) VALUES ({placeholders})", valores
        )
        conn.commit()

        total, n_clientes, n_gestorias = matching.recalcular_todas_las_gestorias(conn)
        flash(f"Convocatoria {id_convocatoria} creada. Cruce recalculado: {total} matches sobre {n_clientes} clientes de {n_gestorias} gestorías.")
        return redirect(url_for("convocatorias"))

    return render_template("convocatoria_form.html", c=None)


@app.route("/convocatorias/<id_convocatoria>/editar", methods=["GET", "POST"])
@admin_requerido
def editar_convocatoria(id_convocatoria):
    conn = get_conn()
    c = conn.execute(
        "SELECT * FROM convocatorias WHERE id_convocatoria = ?", (id_convocatoria,)
    ).fetchone()
    if not c:
        flash("Convocatoria no encontrada.")
        return redirect(url_for("convocatorias"))

    if request.method == "POST":
        datos = _leer_formulario_convocatoria()
        set_clause = ", ".join(f"{campo} = ?" for campo in CAMPOS_CONVOCATORIA)
        valores = [datos[c2] for c2 in CAMPOS_CONVOCATORIA] + [id_convocatoria]
        conn.execute(
            f"UPDATE convocatorias SET {set_clause} WHERE id_convocatoria = ?", valores
        )
        conn.commit()

        total, n_clientes, n_gestorias = matching.recalcular_todas_las_gestorias(conn)
        flash(f"Convocatoria {id_convocatoria} actualizada. Cruce recalculado: {total} matches sobre {n_clientes} clientes de {n_gestorias} gestorías.")
        return redirect(url_for("convocatorias"))

    return render_template("convocatoria_form.html", c=c)


# -------------------------------------------------------------- gestorias (alta de clientes de Subvantia)

def _generar_password(longitud=10):
    alfabeto = string.ascii_letters + string.digits
    return "".join(secrets.choice(alfabeto) for _ in range(longitud))


@app.route("/gestorias")
@admin_requerido
def gestorias():
    conn = get_conn()
    filas = conn.execute(
        """SELECT u.*,
                  (SELECT COUNT(*) FROM clientes c WHERE c.usuario_id = u.id) AS n_clientes,
                  (SELECT COUNT(*) FROM matches m WHERE m.usuario_id = u.id) AS n_matches
           FROM usuarios u
           WHERE u.rol = 'gestoria'
           ORDER BY u.creado_en DESC"""
    ).fetchall()
    return render_template("gestorias.html", gestorias=filas)


@app.route("/gestorias/nueva", methods=["GET", "POST"])
@admin_requerido
def nueva_gestoria():
    conn = get_conn()
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        nombre = request.form.get("nombre", "").strip()
        gestoria = request.form.get("gestoria", "").strip()
        password = request.form.get("password", "").strip()

        if not email or not gestoria:
            flash("El email y el nombre de la gestoría son obligatorios.")
            return redirect(url_for("nueva_gestoria"))

        existe = conn.execute("SELECT 1 FROM usuarios WHERE email = ?", (email,)).fetchone()
        if existe:
            flash(f"Ya existe una cuenta con el email {email}.")
            return redirect(url_for("nueva_gestoria"))

        if not password:
            password = _generar_password()

        conn.execute(
            """INSERT INTO usuarios (email, password_hash, nombre, gestoria, rol, creado_en)
               VALUES (?, ?, ?, ?, 'gestoria', ?)""",
            (email, generate_password_hash(password), nombre, gestoria, dt.datetime.now().isoformat()),
        )
        conn.commit()

        flash(
            f"Cuenta creada para {gestoria}. Envíale estas credenciales para su primer acceso — "
            f"email: {email} · contraseña: {password}"
        )
        return redirect(url_for("gestorias"))

    return render_template("gestoria_form.html")


ESTADOS_SUSCRIPCION = ["activa", "impago", "pausada"]
ESTADOS_SUSCRIPCION_LABEL = {
    "activa": "Activa",
    "impago": "Impago",
    "pausada": "Pausada",
}


@app.route("/gestorias/<int:usuario_id>/estado", methods=["POST"])
@admin_requerido
def cambiar_estado_gestoria(usuario_id):
    conn = get_conn()
    nuevo_estado = request.form.get("estado_suscripcion", "")
    if nuevo_estado not in ESTADOS_SUSCRIPCION:
        flash("Estado no reconocido.")
        return redirect(url_for("gestorias"))

    gestoria = conn.execute(
        "SELECT * FROM usuarios WHERE id = ? AND rol = 'gestoria'", (usuario_id,)
    ).fetchone()
    if not gestoria:
        flash("Gestoría no encontrada.")
        return redirect(url_for("gestorias"))

    conn.execute("UPDATE usuarios SET estado_suscripcion = ? WHERE id = ?", (nuevo_estado, usuario_id))
    conn.commit()

    # Cambiar el estado no borra ni toca su cartera de clientes ni sus matches:
    # si vuelve a ponerse "activa" recupera acceso instantaneo a todo tal cual
    # estaba. Solo afecta a si puede iniciar sesion.
    flash(f"{gestoria['gestoria']} marcada como {ESTADOS_SUSCRIPCION_LABEL[nuevo_estado].lower()}.")
    return redirect(url_for("gestorias"))


# -------------------------------------------------------------------- accesos

@app.route("/accesos")
@admin_requerido
def accesos():
    conn = get_conn()

    demo = conn.execute("SELECT id FROM usuarios WHERE email = ?", (EMAIL_CUENTA_DEMO,)).fetchone()
    demo_id = demo["id"] if demo else None

    resumen_demo = {"hoy": 0, "ultimos_7_dias": 0, "ultimos_30_dias": 0, "total": 0}
    if demo_id:
        ahora = dt.datetime.now()
        filas_demo = conn.execute(
            "SELECT fecha_hora FROM accesos_log WHERE usuario_id = ?", (demo_id,)
        ).fetchall()
        for f in filas_demo:
            try:
                fh = dt.datetime.fromisoformat(f["fecha_hora"])
            except (ValueError, TypeError):
                continue
            dias = (ahora - fh).days
            resumen_demo["total"] += 1
            if dias < 1:
                resumen_demo["hoy"] += 1
            if dias < 7:
                resumen_demo["ultimos_7_dias"] += 1
            if dias < 30:
                resumen_demo["ultimos_30_dias"] += 1

    ultimos_accesos = conn.execute(
        """SELECT a.fecha_hora, a.ip, u.email, u.gestoria, u.rol
           FROM accesos_log a JOIN usuarios u ON u.id = a.usuario_id
           ORDER BY a.fecha_hora DESC
           LIMIT 200"""
    ).fetchall()

    return render_template(
        "accesos.html",
        resumen_demo=resumen_demo,
        ultimos_accesos=ultimos_accesos,
        email_demo=EMAIL_CUENTA_DEMO,
    )


# ----------------------------------------------------------------- matches

@app.route("/matches")
@login_requerido
def matches():
    conn = get_conn()
    uid = session["usuario_id"]
    banda = request.args.get("banda", "")
    q = request.args.get("q", "").strip()

    sql = """SELECT m.*, c.nombre AS cliente_nombre, conv.titulo AS convocatoria_titulo
              FROM matches m
              JOIN clientes c ON c.id = m.cliente_id
              JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
              WHERE m.usuario_id = ?"""
    params = [uid]
    if banda in ("alta", "media", "baja"):
        sql += " AND m.banda = ?"
        params.append(banda)
    if q:
        sql += " AND c.nombre LIKE ?"
        params.append(f"%{q}%")
    sql += " ORDER BY m.puntuacion DESC"

    filas = conn.execute(sql, params).fetchall()

    # expedientes ya creados para saber que botones mostrar
    expedientes_por_match = {
        r["match_id"]: r["estado"]
        for r in conn.execute(
            """SELECT e.match_id, e.estado FROM expedientes e
               JOIN matches m ON m.id = e.match_id WHERE m.usuario_id = ?""",
            (uid,),
        ).fetchall()
    }

    return render_template("matches.html", matches=filas, banda=banda, q=q,
                            expedientes_por_match=expedientes_por_match)


@app.route("/matches/<int:match_id>/email")
@login_requerido
def ver_email(match_id):
    conn = get_conn()
    texto = matching.generar_email(conn, match_id)
    if texto is None:
        flash("Match no encontrado.")
        return redirect(url_for("matches"))
    return render_template("email.html", texto=texto, match_id=match_id)


@app.route("/matches/<int:match_id>/crear_expediente", methods=["POST"])
@login_requerido
def crear_expediente(match_id):
    if es_cuenta_demo():
        flash(MENSAJE_DEMO_SOLO_LECTURA)
        return redirect(request.referrer or url_for("matches"))
    conn = get_conn()
    existente = conn.execute("SELECT id FROM expedientes WHERE match_id = ?", (match_id,)).fetchone()
    if not existente:
        ahora = dt.datetime.now().isoformat()
        conn.execute(
            "INSERT INTO expedientes (match_id, estado, creado_en, actualizado_en) VALUES (?, 'detectada', ?, ?)",
            (match_id, ahora, ahora),
        )
        conn.commit()
        flash("Expediente creado y anadido al kanban.")
    return redirect(request.referrer or url_for("matches"))


# -------------------------------------------------------------- expedientes

@app.route("/expedientes")
@login_requerido
def expedientes():
    conn = get_conn()
    uid = session["usuario_id"]
    filas = conn.execute(
        """SELECT e.*, c.nombre AS cliente_nombre, conv.titulo AS convocatoria_titulo,
                  m.cuantia_estimada
           FROM expedientes e
           JOIN matches m ON m.id = e.match_id
           JOIN clientes c ON c.id = m.cliente_id
           JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ?
           ORDER BY e.actualizado_en DESC""",
        (uid,),
    ).fetchall()

    columnas = {estado: [] for estado in db.ESTADOS_EXPEDIENTE}
    for f in filas:
        columnas.setdefault(f["estado"], []).append(f)

    return render_template("expedientes.html", columnas=columnas,
                            estados=db.ESTADOS_EXPEDIENTE,
                            labels=db.ESTADOS_EXPEDIENTE_LABEL)


@app.route("/expedientes/<int:expediente_id>/avanzar", methods=["POST"])
@login_requerido
def avanzar_expediente(expediente_id):
    if es_cuenta_demo():
        flash(MENSAJE_DEMO_SOLO_LECTURA)
        return redirect(url_for("expedientes"))
    conn = get_conn()
    exp = conn.execute("SELECT * FROM expedientes WHERE id = ?", (expediente_id,)).fetchone()
    if exp:
        estados = db.ESTADOS_EXPEDIENTE
        idx_actual = estados.index(exp["estado"]) if exp["estado"] in estados else 0
        nuevo_estado = request.form.get("estado")
        if nuevo_estado in estados:
            conn.execute(
                "UPDATE expedientes SET estado = ?, actualizado_en = ? WHERE id = ?",
                (nuevo_estado, dt.datetime.now().isoformat(), expediente_id),
            )
            conn.commit()
    return redirect(url_for("expedientes"))


# ------------------------------------------------------------------ export

@app.route("/export/matches.xlsx")
@login_requerido
def exportar_matches():
    conn = get_conn()
    uid = session["usuario_id"]
    filas = conn.execute(
        """SELECT c.nombre AS cliente, c.cif_nif AS cif, conv.titulo AS convocatoria,
                  m.id_convocatoria, m.puntuacion, m.banda, m.motivo, m.cuantia_estimada, m.alerta
           FROM matches m
           JOIN clientes c ON c.id = m.cliente_id
           JOIN convocatorias conv ON conv.id_convocatoria = m.id_convocatoria
           WHERE m.usuario_id = ?
           ORDER BY m.puntuacion DESC""",
        (uid,),
    ).fetchall()

    FONT = "Arial"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados del cruce"
    headers = ["Cliente", "CIF", "Convocatoria", "ID", "Puntuacion", "Banda",
               "Motivo", "Cuantia estimada (EUR)", "Alerta"]
    widths = [26, 12, 34, 10, 11, 9, 40, 16, 30]
    ws.append(headers)
    for c, w in zip(range(1, len(headers) + 1), widths):
        ws.column_dimensions[get_column_letter(c)].width = w
    header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    ws.freeze_panes = "A2"

    banda_fill = {
        "alta": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
        "media": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "baja": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
    }
    for r in filas:
        ws.append([r["cliente"], r["cif"], r["convocatoria"], r["id_convocatoria"],
                    r["puntuacion"], r["banda"], r["motivo"], r["cuantia_estimada"], r["alerta"]])
        row_idx = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = Font(name=FONT, size=10)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row=row_idx, column=6).fill = banda_fill.get(r["banda"], PatternFill())

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="Subvantia_Resultado_Cruce.xlsx",
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# Se inicializa la base de datos al importar el modulo (no solo al ejecutar
# "python app.py" directamente), para que tambien funcione cuando un servidor
# de produccion como gunicorn importa "app" sin pasar por el bloque __main__.
db.init_db()

if __name__ == "__main__":
    modo_debug = os.environ.get("SUBVANTIA_DEBUG", "1") == "1"
    puerto = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=puerto, debug=modo_debug)
